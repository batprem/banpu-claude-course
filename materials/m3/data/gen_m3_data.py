#!/usr/bin/env python3
"""Generate the synthetic M3 Sample Pack (Outlook mailbox + SharePoint Lists + shared Excel).

All people, sites, numbers and emails are fictional (@example.com). Deterministic (seed 303).

Output (small, committed): materials/m3/data/pack/
  outlook/*.eml              test-mailbox emails (project updates, approval requests, meeting summaries)
  outlook/mailbox_index.csv  same emails as one CSV (for IT import or the local fallback)
  sharepoint/Project_Tracker.csv   SharePoint List "Project Tracker"
  sharepoint/Issue_Log.csv         SharePoint List "Issue Log"
  sharepoint/Shared_Capex_Tracker_2569.xlsx  the Excel "many people edit at once" (copy-before-write demo)

Usage: python3 materials/m3/data/gen_m3_data.py
"""
import csv
import datetime as dt
import pathlib
import random
from email.message import EmailMessage
from email.utils import format_datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

OUT = pathlib.Path(__file__).resolve().parent / "pack"
TZ = dt.timezone(dt.timedelta(hours=7))
random.seed(303)

# Same fictional sites as M2 (materials/m2/data/core/ops_monthly_clean.csv)
SITES = {
    "A": "เหมืองสมมติ ภูผาทอง", "B": "เหมืองสมมติ ดอยหินดำ", "C": "เหมืองสมมติ ผาแดง",
    "D": "โรงไฟฟ้าสมมติ ลำน้ำใส", "E": "โรงไฟฟ้าสมมติ ทุ่งทอง", "F": "ฟาร์มโซลาร์สมมติ แสงตะวัน",
}
PEOPLE = {  # obviously fictional
    "somsak": ("สมศักดิ์ ทดสอบ", "somsak.t@example.com", "PM Site A"),
    "wilai": ("วิไลวรรณ สมมติ", "wilaiwan.s@example.com", "Project Control"),
    "kanya": ("กัญญา ตัวอย่าง", "kanya.t@example.com", "PM Site D"),
    "pravit": ("ประวิทย์ จำลอง", "pravit.j@example.com", "วิศวกร Site F"),
    "nok": ("นกน้อย ใจดี", "noknoi.j@example.com", "Finance BP"),
    "chai": ("ชัยวัฒน์ ลองดู", "chaiwat.l@example.com", "Head of Operations"),
    "mali": ("มะลิ ฝึกงาน", "mali.f@example.com", "PM Site B"),
    "trainee": ("ผู้เรียน ทดสอบ", "m3.learner@example.com", "บัญชีทดสอบ (mailbox)"),
}

PROJECTS = [
    # id, name, site, owner, budget MTHB, start, due, status, pct
    ("PRJ-101", "เปลี่ยนสายพานลำเลียง CV-03", "A", "somsak", 48.0, "2026-05-01", "2026-12-15", "On track", 62),
    ("PRJ-102", "ติดตั้งระบบตรวจฝุ่นออนไลน์", "A", "somsak", 12.5, "2026-06-10", "2026-11-30", "At risk", 70),
    ("PRJ-103", "ปรับปรุงถนนขนส่งในเหมือง", "B", "mali", 31.0, "2026-03-01", "2027-02-28", "On track", 40),
    ("PRJ-104", "ระบบบำบัดน้ำทิ้งเฟส 2", "C", "mali", 26.4, "2026-04-15", "2026-12-31", "Delayed", 35),
    ("PRJ-105", "Overhaul หม้อไอน้ำ Unit 2", "D", "kanya", 95.0, "2026-08-01", "2026-11-20", "On track", 55),
    ("PRJ-106", "เปลี่ยน DCS โรงไฟฟ้า", "E", "kanya", 72.3, "2026-02-01", "2027-03-31", "At risk", 30),
    ("PRJ-107", "เปลี่ยน inverter ชุด 4–6", "F", "pravit", 18.9, "2026-07-01", "2026-12-10", "On track", 48),
    ("PRJ-108", "ติดตั้งแบตเตอรี่สำรอง (BESS) นำร่อง", "F", "pravit", 64.0, "2026-09-01", "2027-06-30", "Not started", 5),
    ("PRJ-109", "Digital permit-to-work", "B", "wilai", 6.8, "2026-05-20", "2026-12-20", "On track", 58),
    ("PRJ-110", "ปรับปรุงห้องควบคุม Unit 1", "D", "kanya", 22.0, "2026-06-01", "2026-12-31", "On track", 50),
    ("PRJ-111", "ระบบพยากรณ์ผลผลิตโซลาร์", "F", "pravit", 4.2, "2026-08-15", "2027-01-31", "On track", 25),
    ("PRJ-112", "ซ่อมบำรุงใหญ่รถขุด EX-07", "C", "mali", 15.6, "2026-09-10", "2026-11-25", "At risk", 60),
]

ISSUE_TEMPLATES = [
    ("อะไหล่ล่าช้าจาก supplier", "Supply", "High"),
    ("ผลทดสอบไม่ผ่านรอบแรก", "Quality", "Medium"),
    ("ฝนตกหนัก หยุดงานหน้างาน", "Weather", "Medium"),
    ("รอใบอนุญาตจากหน่วยงานรัฐ", "Permit", "High"),
    ("ค่าใช้จ่ายเกินงบช่วง", "Cost", "High"),
    ("ผู้รับเหมาขาดแรงงาน", "Resource", "Medium"),
    ("แบบก่อสร้างต้องแก้ไข", "Design", "Low"),
    ("near-miss ระหว่างยกของ", "Safety", "High"),
    ("ข้อมูล SCADA ไม่ส่งเข้าระบบ", "IT/OT", "Medium"),
    ("รอผลอนุมัติ change order", "Approval", "Medium"),
]


def person(key):
    return PEOPLE[key]


def write_csv(path, rows, header):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def project_tracker():
    rows = []
    for pid, name, site, owner, budget, start, due, status, pct in PROJECTS:
        spent = round(budget * pct / 100 * random.uniform(0.85, 1.2), 1)
        rows.append([pid, name, f"Site {site}", SITES[site], person(owner)[0], person(owner)[1],
                     budget, spent, start, due, status, pct, "2026-10-24"])
    write_csv(OUT / "sharepoint" / "Project_Tracker.csv", rows,
              ["ProjectID", "Title", "Site", "SiteName", "Owner", "OwnerEmail", "Budget_MTHB",
               "Spent_MTHB", "StartDate", "DueDate", "Status", "PercentComplete", "Modified"])
    return rows


def issue_log():
    rows, n = [], 1
    base = dt.date(2026, 7, 1)
    for pid, name, site, owner, *_ in PROJECTS:
        for _ in range(random.randint(3, 7)):
            title, cat, sev = random.choice(ISSUE_TEMPLATES)
            opened = base + dt.timedelta(days=random.randint(0, 115))
            closed = random.random() < 0.55
            closed_on = (opened + dt.timedelta(days=random.randint(2, 30))) if closed else None
            if closed_on and closed_on > dt.date(2026, 10, 28):
                closed_on, closed = None, False
            rows.append([f"ISS-{n:03d}", pid, f"Site {site}", title, cat, sev,
                         "Closed" if closed else random.choice(["Open", "Open", "In progress"]),
                         person(owner)[0], opened.isoformat(), closed_on.isoformat() if closed_on else ""])
            n += 1
    # Planted issues the emails refer to (cross-source search targets)
    rows += [
        [f"ISS-{n:03d}", "PRJ-101", "Site A", "ลูกกลิ้งสายพานล็อตใหม่ส่งไม่ครบ 40 ชิ้น", "Supply", "High",
         "Open", person("somsak")[0], "2026-10-20", ""],
        [f"ISS-{n+1:03d}", "PRJ-105", "Site D", "พบรอยร้าวท่อ superheater เพิ่ม 2 จุด", "Quality", "High",
         "In progress", person("kanya")[0], "2026-10-22", ""],
    ]
    write_csv(OUT / "sharepoint" / "Issue_Log.csv", rows,
              ["IssueID", "ProjectID", "Site", "Title", "Category", "Severity", "Status", "Owner",
               "OpenedDate", "ClosedDate"])
    return rows


EMAILS = [
    # (days before 2026-10-28, from, subject, category, body)
    (1, "somsak", "[PRJ-101] อัปเดตสัปดาห์ 43: ลูกกลิ้งสายพานส่งไม่ครบ", "project-update",
     "เรียนทีม\n\nสัปดาห์นี้ PRJ-101 เปลี่ยนสายพาน CV-03 คืบหน้า 62% ตามแผน\n"
     "แต่ supplier ส่งลูกกลิ้งล็อตใหม่ไม่ครบ 40 ชิ้น (บันทึกใน Issue Log แล้ว)\n"
     "ถ้าไม่ได้ของภายใน 7 พ.ย. อาจเลื่อนเสร็จจาก 15 ธ.ค. เป็นปลายเดือน ธ.ค.\n"
     "ตอนนี้ใน Project Tracker ยังเป็น On track อยู่ ขอทบทวนวันศุกร์\n\nสมศักดิ์"),
    (2, "kanya", "[PRJ-105] ขออนุมัติงบเพิ่ม 3.8 ล้านบาท (Overhaul Unit 2)", "approval-request",
     "เรียนคุณชัยวัฒน์\n\nระหว่าง overhaul หม้อไอน้ำ Unit 2 พบรอยร้าวท่อ superheater เพิ่ม 2 จุด\n"
     "ขออนุมัติงบเพิ่ม 3.8 ล้านบาท (จาก 95.0 เป็น 98.8 ล้านบาท) เพื่อเปลี่ยนท่อช่วงนี้ทันที\n"
     "ถ้าไม่อนุมัติภายใน 31 ต.ค. ต้องเลื่อน start-up ออกไป ~10 วัน\n"
     "รายละเอียดอยู่ในไฟล์ Shared_Capex_Tracker_2569.xlsx แท็บ Site D\n\nกัญญา"),
    (3, "wilai", "สรุปประชุม Project Review ประจำเดือน ต.ค. 2569", "meeting-summary",
     "สรุปประชุม Project Review (ข้อมูลสมมติ)\n"
     "ผู้เข้าประชุม: คุณชัยวัฒน์, คุณสมศักดิ์, คุณกัญญา, คุณประวิทย์, คุณมะลิ, คุณนกน้อย\n\n"
     "1) PRJ-104 บำบัดน้ำทิ้งเฟส 2 ล่าช้า รอใบอนุญาต คุณมะลิติดตาม\n"
     "2) PRJ-106 DCS: เสี่ยงงบเกิน ให้ Finance ทบทวนตัวเลขใน Capex Tracker\n"
     "3) PRJ-108 BESS ยังไม่เริ่ม รอผลศึกษาความเป็นไปได้\n"
     "4) ขอให้ทุก PM อัปเดต Capex Tracker ก่อนวันศุกร์ที่ 30 ต.ค. **ห้ามแก้แท็บของไซต์อื่น**\n\n"
     "Action: Project Control รวบรวมสถานะส่งผู้บริหารวันจันทร์"),
    (4, "pravit", "[PRJ-107] inverter ชุด 4 ติดตั้งเสร็จ ทดสอบผ่าน", "project-update",
     "inverter ชุด 4 ติดตั้งและทดสอบผ่านแล้ว ชุด 5–6 ตามแผนเดือน พ.ย.\nสถานะ On track 48%\n\nประวิทย์"),
    (5, "mali", "[PRJ-112] รถขุด EX-07 รออะไหล่ไฮดรอลิก", "project-update",
     "อะไหล่ปั๊มไฮดรอลิกยังไม่มา คาดว่าได้ 5 พ.ย. งานซ่อมอาจเลื่อนจาก 25 พ.ย. เป็น 2 ธ.ค.\nมะลิ"),
    (6, "nok", "ขอทบทวนตัวเลข Spent ใน Capex Tracker ก่อนปิดเดือน", "approval-request",
     "เรียน PM ทุกไซต์\nFinance พบว่าบางแท็บใน Capex Tracker ตัวเลข Spent ไม่ตรงกับระบบบัญชี\n"
     "ขอให้ตรวจแท็บของตัวเองและยืนยันภายใน 30 ต.ค.\nนกน้อย (Finance BP)"),
    (7, "chai", "RE: [PRJ-106] ขอเลื่อน go-live DCS", "approval-request",
     "อนุมัติให้เลื่อน go-live DCS จาก ม.ค. เป็น มี.ค. 2570 ตามที่เสนอ\nแต่ขอแผนลดความเสี่ยงงบเกินภายในสัปดาห์หน้า\nชัยวัฒน์"),
    (8, "kanya", "[PRJ-110] ปรับปรุงห้องควบคุม: ส่งแบบรอบ 2", "project-update",
     "ส่งแบบห้องควบคุมรอบ 2 ให้ทีมความปลอดภัยตรวจ คาดว่าได้ผล 4 พ.ย.\nกัญญา"),
    (9, "wilai", "สรุปประชุม Safety Stand-down Site A", "meeting-summary",
     "สรุป: near-miss ระหว่างยกลูกกลิ้งสายพาน (PRJ-101) ไม่มีผู้บาดเจ็บ\n"
     "มาตรการ: ทบทวน lifting plan ทุกงาน, เพิ่มผู้ให้สัญญาณ 1 คน\nบันทึกใน Issue Log หมวด Safety"),
    (10, "mali", "[PRJ-104] ใบอนุญาตระบายน้ำยังไม่ออก", "project-update",
     "หน่วยงานรัฐขอเอกสารเพิ่ม 2 รายการ คาดว่าใบอนุญาตออกกลาง พ.ย. สถานะ Delayed\nมะลิ"),
    (11, "pravit", "ขออนุมัติจ้างที่ปรึกษาศึกษา BESS (PRJ-108)", "approval-request",
     "ขออนุมัติวงเงิน 1.2 ล้านบาท จ้างที่ปรึกษาศึกษาความเป็นไปได้ BESS 8 สัปดาห์\nประวิทย์"),
    (12, "somsak", "[PRJ-102] ระบบตรวจฝุ่น: เซ็นเซอร์ 2 ตัวค่าเพี้ยน", "project-update",
     "เซ็นเซอร์ฝุ่น 2 จาก 8 ตัวอ่านค่าเพี้ยน ต้อง calibrate ใหม่ สถานะ At risk\nสมศักดิ์"),
    (14, "wilai", "สรุปประชุมทีม PM รายสัปดาห์ (สัปดาห์ 42)", "meeting-summary",
     "1) ทุกโครงการอัปเดต Issue Log ภายในวันพฤหัส\n2) Capex Tracker ใช้ไฟล์เดียวบน SharePoint ห้ามส่งไฟล์แนบเวอร์ชันส่วนตัว\n"
     "3) ถ้าจะให้ AI ช่วยแก้ไฟล์ ให้ทำสำเนาก่อนเสมอ (copy-before-write)"),
    (15, "nok", "FYI: ปิดงบไตรมาส 3 เรียบร้อย", "fyi",
     "ปิดงบไตรมาส 3 เรียบร้อย ตัวเลข Spent รวมทุกโครงการอัปเดตใน Capex Tracker แล้ว\nนกน้อย"),
    (16, "mali", "[PRJ-103] ถนนขนส่ง: เทคอนกรีตช่วงที่ 2 เสร็จ", "project-update",
     "ช่วงที่ 2 เสร็จ ความคืบหน้ารวม 40% ตามแผน\nมะลิ"),
    (18, "chai", "แจ้งกำหนดส่งรายงานสถานะโครงการให้ผู้บริหาร", "fyi",
     "รายงานสถานะโครงการทุกไซต์ ส่งทุกวันจันทร์ก่อน 10:00 หนึ่งหน้า สรุปความเสี่ยงสูงสุด 3 เรื่อง\nชัยวัฒน์"),
    (20, "kanya", "[PRJ-106] DCS: FAT รอบแรกพบ 14 รายการ", "project-update",
     "Factory Acceptance Test รอบแรกพบประเด็น 14 รายการ ปิดได้ 9 รายการ เหลือ 5\nกัญญา"),
    (22, "pravit", "สรุปประชุมกับ supplier inverter", "meeting-summary",
     "supplier ยืนยันส่งชุด 5–6 วันที่ 12 พ.ย. ทีมติดตั้งพร้อม\nประวิทย์"),
    (25, "wilai", "แบบฟอร์ม Issue Log ใหม่บน SharePoint", "fyi",
     "ย้าย Issue Log ไปเป็น SharePoint List แล้ว ใช้คอลัมน์ Severity/Status ตาม dropdown\nวิไลวรรณ"),
    (27, "somsak", "ขออนุมัติ OT ทีมติดตั้งสายพาน 2 สัปดาห์", "approval-request",
     "ขออนุมัติ OT 2 สัปดาห์ช่วงเปลี่ยนสายพาน CV-03 ประมาณ 0.6 ล้านบาท\nสมศักดิ์"),
]


def mailbox():
    now = dt.datetime(2026, 10, 28, 9, 0, tzinfo=TZ)
    rows = []
    folder = OUT / "outlook"
    folder.mkdir(parents=True, exist_ok=True)
    for i, (days, frm, subject, cat, body) in enumerate(EMAILS, 1):
        when = now - dt.timedelta(days=days, hours=random.randint(0, 8), minutes=random.randint(0, 59))
        name, addr, role = person(frm)
        msg = EmailMessage()
        msg["From"] = f"{name} <{addr}>"
        msg["To"] = f"{PEOPLE['trainee'][0]} <{PEOPLE['trainee'][1]}>"
        msg["Cc"] = "pm-team@example.com"
        msg["Subject"] = subject
        msg["Date"] = format_datetime(when)
        msg["Message-ID"] = f"<m3-{i:03d}@example.com>"
        msg["X-M3-Category"] = cat
        msg.set_content(body + f"\n\n--\n{name} · {role}\n(ข้อมูลสมมติเพื่อการฝึกอบรม)\n")
        (folder / f"{i:02d}_{cat}.eml").write_bytes(bytes(msg))
        rows.append([f"m3-{i:03d}", when.isoformat(), name, addr, PEOPLE["trainee"][1], subject, cat,
                     body.replace("\n", " ")])
    write_csv(folder / "mailbox_index.csv", rows,
              ["MessageID", "Received", "FromName", "FromEmail", "To", "Subject", "Category", "Body"])


def shared_excel(tracker):
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="2E86AB")
    warn = PatternFill("solid", fgColor="F2A541")
    ws = wb.active
    ws.title = "อ่านก่อน"
    for r in [
        ["Shared_Capex_Tracker_2569.xlsx (ข้อมูลสมมติ)"],
        ["ไฟล์นี้หลายคนแก้พร้อมกันบน SharePoint: PM แต่ละไซต์แก้เฉพาะแท็บของตัวเอง, Finance แก้คอลัมน์ Spent_Confirmed"],
        ["แท็บ Summary คำนวณด้วยสูตร ห้ามพิมพ์ทับ"],
        ["กฎ: ก่อนให้ Claude แก้ไฟล์นี้ ให้ทำสำเนาก่อนเสมอ (copy-before-write) แล้วตรวจก่อนบันทึก"],
    ]:
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=14)
    ws["A4"].fill = warn
    ws.column_dimensions["A"].width = 110

    cols = ["ProjectID", "Title", "Budget_MTHB", "Spent_PM", "Spent_Confirmed", "Forecast_MTHB",
            "Variance_MTHB", "Status", "LastEditedBy", "Note"]
    dv = DataValidation(type="list", formula1='"On track,At risk,Delayed,Not started,Done"')
    site_refs = []
    for code in SITES:
        rows = [p for p in tracker if p[2] == f"Site {code}"]
        if not rows:
            continue
        sh = wb.create_sheet(f"Site {code}")
        sh.append(cols)
        for c in sh[1]:
            c.font, c.fill = head, fill
        sh.add_data_validation(dv)
        for p in rows:
            r = sh.max_row + 1
            forecast = round(p[6] * random.uniform(0.97, 1.12), 1)
            if p[0] == "PRJ-105":  # workshop: learners update this to 98.8 per the approval email
                forecast = p[6]
            sh.append([p[0], p[1], p[6], p[7], round(p[7] * random.uniform(0.95, 1.0), 1), forecast,
                       f"=F{r}-C{r}", p[10], p[4], ""])
            dv.add(f"H{r}")
        last = sh.max_row
        sh.append(["รวม", "", f"=SUM(C2:C{last})", f"=SUM(D2:D{last})", f"=SUM(E2:E{last})",
                   f"=SUM(F2:F{last})", f"=SUM(G2:G{last})"])
        for c in sh[sh.max_row]:
            c.font = Font(bold=True)
        for col, w in zip("ABCDEFGHIJ", [10, 36, 13, 11, 16, 15, 15, 12, 18, 24]):
            sh.column_dimensions[col].width = w
        site_refs.append((code, sh.title, sh.max_row))

    sm = wb.create_sheet("Summary", 1)
    sm.append(["Site", "SiteName", "Budget_MTHB", "Spent_Confirmed", "Forecast_MTHB", "Variance_MTHB"])
    for c in sm[1]:
        c.font, c.fill = head, fill
    for code, title, tr in site_refs:
        q = f"'{title}'"
        sm.append([title, SITES[code], f"={q}!C{tr}", f"={q}!E{tr}", f"={q}!F{tr}", f"={q}!G{tr}"])
    n = sm.max_row
    sm.append(["รวมทุกไซต์", "", f"=SUM(C2:C{n})", f"=SUM(D2:D{n})", f"=SUM(E2:E{n})", f"=SUM(F2:F{n})"])
    for c in sm[sm.max_row]:
        c.font = Font(bold=True)
    for col, w in zip("ABCDEF", [12, 28, 14, 16, 15, 15]):
        sm.column_dimensions[col].width = w

    log = wb.create_sheet("Change_Log")
    log.append(["When", "Who", "Sheet", "What"])
    for c in log[1]:
        c.font, c.fill = head, fill
    log.append(["2026-10-26 16:10", "กัญญา ตัวอย่าง", "Site D", "อัปเดต Forecast PRJ-105 หลังพบรอยร้าว"])
    log.append(["2026-10-27 09:45", "นกน้อย ใจดี", "Site E", "ยืนยัน Spent_Confirmed ไตรมาส 3"])
    log.append(["2026-10-27 14:20", "สมศักดิ์ ทดสอบ", "Site A", "อัปเดต Spent_PM PRJ-101/102"])
    log.column_dimensions["D"].width = 50
    (OUT / "sharepoint").mkdir(parents=True, exist_ok=True)
    wb.save(OUT / "sharepoint" / "Shared_Capex_Tracker_2569.xlsx")


def main():
    tracker = project_tracker()
    issue_log()
    mailbox()
    shared_excel(tracker)
    total = sum(f.stat().st_size for f in OUT.rglob("*") if f.is_file())
    print(f"{OUT}: {sum(1 for f in OUT.rglob('*') if f.is_file())} files, {total/1024:.0f} KB")


if __name__ == "__main__":
    main()
