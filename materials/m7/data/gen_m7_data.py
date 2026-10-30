#!/usr/bin/env python3
"""Generate the M7 Capstone data pack (synthetic, small, committed).

Five use cases standing in for the ones Hi-I will send. Every name, site,
number and email is invented. Run from the repo root:

    python3 materials/m7/data/gen_m7_data.py

Writes materials/m7/data/uc1_* … uc5_* (commit these) and one ZIP per use
case in materials/m7/data/out/ (gitignored, hand out to teams).
"""
import csv
import datetime as dt
import pathlib
import random
import zipfile

import docx
import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent
random.seed(7)

SITES = {
    "A": "เหมืองสมมติ ภูผาทอง",
    "C": "โรงไฟฟ้าสมมติ ลำน้ำใส",
    "F": "ฟาร์มโซลาร์สมมติ แสงตะวัน",
    "G": "โรงไฟฟ้าสมมติ ธาราทิพย์",
}
NOTE = "ข้อมูลสมมติเพื่อการอบรมเท่านั้น (synthetic) ไม่ใช่ข้อมูลจริงของบ้านปู"


def write_csv(path, header, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def write_text(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.strip() + "\n", encoding="utf-8")


# ---------------------------------------------------------------- UC1
CONTRACTS = [
    {
        "id": "CT-2569-014", "file": "CT-2569-014_conveyor_spares.docx",
        "title": "สัญญาจัดซื้ออะไหล่สายพานลำเลียง (Conveyor Belt Spare Parts)",
        "supplier": "บริษัท สายพานรุ่งเรือง จำกัด (สมมติ)", "site": "A",
        "value": 18_400_000, "start": "2569-01-01", "end": "2570-12-31",
        "clauses": [
            ("ขอบเขตงาน", "ผู้ขายจัดส่งอะไหล่สายพานลำเลียง ลูกกลิ้ง และชุดซ่อม ตามใบสั่งซื้อรายไตรมาส ให้ เหมืองสมมติ ภูผาทอง (Site A)"),
            ("มูลค่าและการชำระเงิน", "มูลค่ารวมไม่เกิน 18,400,000 บาท (ไม่รวม VAT) ชำระภายใน 45 วันหลังรับของและตรวจรับครบ"),
            ("ระยะเวลาส่งมอบ", "ส่งมอบภายใน 21 วันนับจากวันที่ออกใบสั่งซื้อ กรณีเร่งด่วนภายใน 7 วัน คิดค่าขนส่งเพิ่มตามจริง"),
            ("ค่าปรับ", "ส่งมอบล่าช้า ปรับวันละ 0.1% ของมูลค่าใบสั่งซื้อ สูงสุดไม่เกิน 10%"),
            ("การรับประกัน", "รับประกันอะไหล่ 12 เดือนนับจากวันติดตั้ง หรือ 18 เดือนนับจากวันส่งมอบ แล้วแต่อย่างใดถึงก่อน"),
            ("การปรับราคา", "ปรับราคาได้ปีละ 1 ครั้ง ไม่เกิน 5% อ้างอิงดัชนีราคาเหล็ก โดยต้องแจ้งล่วงหน้า 60 วัน"),
            ("การต่ออายุ", "ต่ออายุอัตโนมัติคราวละ 1 ปี หากไม่มีฝ่ายใดแจ้งยกเลิกเป็นหนังสือล่วงหน้า 90 วันก่อนสิ้นสุดสัญญา"),
            ("การยกเลิกสัญญา", "ผู้ซื้อยกเลิกได้หากผู้ขายส่งมอบล่าช้าเกิน 30 วัน รวม 3 ครั้งในรอบ 12 เดือน"),
            ("ความปลอดภัย", "พนักงานของผู้ขายที่เข้าพื้นที่เหมืองต้องผ่านการอบรมความปลอดภัยของ Site และสวม PPE ครบ"),
        ],
    },
    {
        "id": "CT-2569-022", "file": "CT-2569-022_camp_catering.docx",
        "title": "สัญญาจ้างบริการอาหารและแม่บ้านแคมป์พนักงาน",
        "supplier": "บริษัท ครัวภูเขา เซอร์วิส จำกัด (สมมติ)", "site": "A",
        "value": 9_600_000, "start": "2569-04-01", "end": "2570-03-31",
        "clauses": [
            ("ขอบเขตงาน", "บริการอาหาร 3 มื้อ และทำความสะอาดแคมป์พนักงาน 120 ห้อง ที่ เหมืองสมมติ ภูผาทอง"),
            ("มูลค่าและการชำระเงิน", "ค่าบริการเหมาจ่ายเดือนละ 800,000 บาท ชำระภายใน 30 วันหลังได้รับใบแจ้งหนี้"),
            ("ระดับบริการ (SLA)", "คะแนนความพึงพอใจเฉลี่ยรายเดือนไม่ต่ำกว่า 4.0/5 และผ่านตรวจสุขอนามัยทุกไตรมาส"),
            ("ค่าปรับ", "หาก SLA ต่ำกว่าเกณฑ์ 2 เดือนติดต่อกัน หักค่าบริการ 5% ของเดือนถัดไป"),
            ("การต่ออายุ", "ไม่ต่ออายุอัตโนมัติ ต้องทำบันทึกต่ออายุเป็นหนังสือ"),
            ("การยกเลิกสัญญา", "ฝ่ายใดฝ่ายหนึ่งยกเลิกได้โดยแจ้งล่วงหน้า 60 วัน"),
            ("ข้อมูลส่วนบุคคล", "ผู้ให้บริการต้องไม่นำรายชื่อหรือข้อมูลพนักงานในแคมป์ไปใช้เพื่อวัตถุประสงค์อื่น"),
        ],
    },
    {
        "id": "CT-2569-031", "file": "CT-2569-031_solar_cleaning.docx",
        "title": "สัญญาจ้างล้างแผงโซลาร์และตัดหญ้า",
        "supplier": "ห้างหุ้นส่วนจำกัด แผงใสสะอาด (สมมติ)", "site": "F",
        "value": 3_150_000, "start": "2569-02-01", "end": "2570-01-31",
        "clauses": [
            ("ขอบเขตงาน", "ล้างแผงโซลาร์ 42,000 แผง เดือนละ 1 รอบ (ฤดูแล้งเดือนละ 2 รอบ) และตัดหญ้ารอบแผงทุก 6 สัปดาห์ ที่ ฟาร์มโซลาร์สมมติ แสงตะวัน (Site F)"),
            ("มูลค่าและการชำระเงิน", "ค่าจ้างรวม 3,150,000 บาท แบ่งจ่ายรายเดือนตามรอบงานที่ตรวจรับแล้ว ชำระภายใน 30 วัน"),
            ("มาตรฐานงาน", "หลังล้าง ค่า soiling loss ต้องต่ำกว่า 2% ตามการวัดของผู้ว่าจ้าง ใช้น้ำไม่เกิน 1.5 ลิตรต่อแผง"),
            ("ค่าปรับ", "งานไม่ผ่านมาตรฐานต้องล้างซ้ำภายใน 5 วันโดยไม่คิดเงิน หากเกิน ปรับวันละ 5,000 บาท"),
            ("ประกันภัย", "ผู้รับจ้างทำประกันความเสียหายต่อทรัพย์สินบุคคลภายนอก วงเงินไม่ต่ำกว่า 5,000,000 บาท"),
            ("การต่ออายุ", "ต่ออายุอัตโนมัติ 1 ปี หากไม่แจ้งยกเลิกล่วงหน้า 30 วัน"),
            ("การยกเลิกสัญญา", "ผู้ว่าจ้างยกเลิกได้ทันทีหากผู้รับจ้างทำแผงเสียหายจากความประมาทซ้ำ 2 ครั้ง"),
        ],
    },
]

REGISTER_EXTRA = [
    ("CT-2568-087", "บริษัท น้ำมันหล่อลื่นไทยสมมติ จำกัด", "A", 12_700_000, "2568-07-01", "2569-12-31", "อัตโนมัติ 1 ปี / แจ้ง 60 วัน"),
    ("CT-2568-091", "บริษัท ยางล้อเหมืองสมมติ จำกัด", "A", 25_300_000, "2568-10-01", "2570-09-30", "ไม่ต่ออัตโนมัติ"),
    ("CT-2569-003", "บริษัท ตรวจวัดอากาศสมมติ จำกัด", "C", 1_850_000, "2569-01-15", "2569-12-31", "ไม่ต่ออัตโนมัติ"),
    ("CT-2569-008", "บริษัท ซ่อมบำรุงหม้อไอน้ำสมมติ จำกัด", "C", 31_000_000, "2569-03-01", "2571-02-28", "อัตโนมัติ 1 ปี / แจ้ง 90 วัน"),
    ("CT-2569-017", "บริษัท รักษาความปลอดภัยสมมติ จำกัด", "G", 4_200_000, "2569-01-01", "2569-12-31", "อัตโนมัติ 1 ปี / แจ้ง 30 วัน"),
    ("CT-2569-025", "บริษัท อินเวอร์เตอร์เซอร์วิสสมมติ จำกัด", "F", 6_750_000, "2569-05-01", "2571-04-30", "ไม่ต่ออัตโนมัติ"),
    ("CT-2569-029", "บริษัท ขนส่งถ่านหินสมมติ จำกัด", "A", 48_900_000, "2569-02-01", "2570-01-31", "อัตโนมัติ 1 ปี / แจ้ง 90 วัน"),
    ("CT-2569-033", "บริษัท ระบบบำบัดน้ำสมมติ จำกัด", "G", 2_400_000, "2569-06-01", "2570-05-31", "อัตโนมัติ 1 ปี / แจ้ง 60 วัน"),
    ("CT-2569-036", "บริษัท ไอทีซัพพอร์ตสมมติ จำกัด", "C", 980_000, "2569-07-01", "2570-06-30", "ไม่ต่ออัตโนมัติ"),
]


def gen_uc1():
    base = ROOT / "uc1_contract_summary"
    for c in CONTRACTS:
        d = docx.Document()
        d.add_heading(c["title"], 0)
        d.add_paragraph(f"เลขที่สัญญา {c['id']} · {NOTE}")
        d.add_paragraph(
            f"คู่สัญญา: ผู้ซื้อ/ผู้ว่าจ้าง บริษัทในกลุ่มธุรกิจพลังงานสมมติ (Site {c['site']} · {SITES[c['site']]}) "
            f"และ {c['supplier']} · ระยะเวลา {c['start']} ถึง {c['end']}")
        for i, (h, body) in enumerate(c["clauses"], 1):
            d.add_heading(f"ข้อ {i} {h}", level=2)
            d.add_paragraph(body)
            d.add_paragraph("คู่สัญญาตกลงให้ข้อนี้ใช้ร่วมกับเงื่อนไขทั่วไปของการจัดซื้อจัดจ้างฉบับปี 2568 "
                            "หากข้อความขัดกัน ให้ถือตามสัญญาฉบับนี้")
        d.add_paragraph("ผู้ติดต่อฝ่ายจัดซื้อ: คุณสมมติ ใจดี · procurement.demo@example.com")
        path = base / "contracts" / c["file"]
        path.parent.mkdir(parents=True, exist_ok=True)
        d.save(path)
    rows = [(c["id"], c["supplier"], c["site"], c["value"], c["start"], c["end"],
             {"CT-2569-014": "อัตโนมัติ 1 ปี / แจ้ง 90 วัน", "CT-2569-022": "ไม่ต่ออัตโนมัติ",
              "CT-2569-031": "อัตโนมัติ 1 ปี / แจ้ง 30 วัน"}[c["id"]]) for c in CONTRACTS]
    rows += REGISTER_EXTRA
    rows.sort()
    write_csv(base / "contract_register.csv",
              ["contract_id", "supplier", "site", "value_thb", "start_be", "end_be", "renewal"], rows)


# ---------------------------------------------------------------- UC2
INC_TYPES = [("Near miss", 0.45), ("First aid", 0.25), ("Equipment damage", 0.12),
             ("Environmental", 0.08), ("Lost time injury", 0.03), ("Unsafe condition", 0.07)]
AREAS = {"A": ["หน้าเหมือง", "โรงซ่อมบำรุง", "ลานกองถ่าน", "ถนนขนส่ง"],
         "C": ["หม้อไอน้ำ", "ห้องควบคุม", "ลานกองเชื้อเพลิง", "ระบบหล่อเย็น"],
         "F": ["แถวแผง 1–40", "แถวแผง 41–80", "ห้องอินเวอร์เตอร์", "ถนนภายใน"],
         "G": ["กังหัน", "หม้อแปลง", "ระบบบำบัดน้ำ", "โกดังสารเคมี"]}
CAUSES = ["ไม่ใช้ PPE ครบ", "พื้นลื่น", "อุปกรณ์ชำรุด", "สื่อสารไม่ชัดระหว่างกะ",
          "ขับรถเร็วเกินกำหนด", "ไม่ล็อกเอาต์ก่อนซ่อม", "ฝนตกหนัก", "จัดเก็บของไม่เป็นระเบียบ"]
REPORTERS = ["คุณสมมติ ระวังภัย", "คุณทดสอบ ปลอดภัยดี", "คุณตัวอย่าง ใส่ใจ", "คุณจำลอง มั่นคง"]


def gen_uc2():
    base = ROOT / "uc2_incident_weekly"
    start = dt.date(2026, 8, 3)  # Monday, 8 weeks (Aug–Sep 2569)
    rows, n = [], 0
    for day in range(56):
        date = start + dt.timedelta(days=day)
        for _ in range(random.choice([1, 2, 2, 3, 3, 4])):
            n += 1
            site = random.choice("AAACCFG")
            r, acc, typ = random.random(), 0, INC_TYPES[0][0]
            for t, p in INC_TYPES:
                acc += p
                if r <= acc:
                    typ = t
                    break
            sev = {"Near miss": 1, "Unsafe condition": 1, "First aid": 2, "Environmental": 2,
                   "Equipment damage": 3, "Lost time injury": 4}[typ]
            status = "Closed" if day < 42 or random.random() < 0.4 else random.choice(["Open", "In progress"])
            rows.append([f"INC-{n:04d}", date.isoformat(), f"Site {site}", random.choice(AREAS[site]),
                         typ, sev, random.choice(CAUSES), status,
                         random.choice(REPORTERS) if site in "AC" else "คุณสมมติ ระวังภัย",
                         "" if status == "Closed" else (date + dt.timedelta(days=14)).isoformat()])
    # two deliberate data issues for the HITL check
    rows[17][4] = "near miss"          # inconsistent casing
    rows[60][1] = "2026-13-02"         # impossible date
    write_csv(base / "incident_log.csv",
              ["incident_id", "date", "site", "area", "type", "severity_1to4", "root_cause",
               "status", "reported_by", "due_date"], rows)
    emails = {
        "01_site_A_weekly_note.eml": ("hse.siteA@example.com", "อัปเดตความปลอดภัย Site A สัปดาห์ที่ 8",
            "สัปดาห์นี้มี near miss บนถนนขนส่ง 3 ครั้ง สาเหตุหลักขับรถเร็วช่วงฝนตก\n"
            "เสนอให้ลดความเร็วสูงสุดเหลือ 30 กม./ชม. ช่วงฝนตก รอผู้จัดการ Site อนุมัติ\n"
            "เหตุ Lost time injury เมื่อวันที่ 20 ก.ย. ปิดเคสแล้ว พนักงานกลับมาทำงานปกติ (ชื่อพนักงานไม่ระบุ)"),
        "02_site_C_boiler.eml": ("hse.siteC@example.com", "Site C: Equipment damage ที่หม้อไอน้ำ",
            "ท่อไอน้ำรั่วเล็กน้อยที่หม้อไอน้ำหมายเลข 2 ไม่มีผู้บาดเจ็บ\n"
            "ผู้รับเหมาซ่อมบำรุง (สัญญา CT-2569-008 สมมติ) จะเข้าซ่อมสัปดาห์หน้า\n"
            "ขอให้รายงานประจำสัปดาห์ระบุเป็นประเด็นติดตาม"),
        "03_hse_manager_request.eml": ("hse.manager@example.com", "ขอรายงาน HSE ประจำสัปดาห์ภายในวันจันทร์ 10:00",
            "รบกวนสรุปเหตุการณ์ทั้งหมดของสัปดาห์ล่าสุดเทียบสัปดาห์ก่อน แยกตาม Site และประเภท\n"
            "ไฮไลต์เหตุที่ severity 3 ขึ้นไป และรายการที่ยังเปิดอยู่เกินกำหนด\n"
            "ขอเป็นร่างอีเมลให้ผมตรวจก่อนส่งผู้บริหาร ห้ามส่งออกเอง"),
    }
    for fname, (frm, subj, body) in emails.items():
        write_text(base / "emails" / fname,
                   f"From: {frm}\nTo: capstone.team@example.com\nSubject: {subj}\n"
                   f"Date: Fri, 25 Sep 2026 16:30:00 +0700\nContent-Type: text/plain; charset=utf-8\n\n"
                   f"{body}\n\n-- {NOTE}")


# ---------------------------------------------------------------- UC3
MONTHS_TH = ["ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค."]


def gen_uc3():
    base = ROOT / "uc3_esg_consolidation"
    # Site A (mine): xlsx, one sheet per topic, electricity in kWh, diesel in litres
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Energy"
    ws.append(["เดือน (2568)", "ไฟฟ้า (kWh)", "ดีเซล (ลิตร)"])
    for i, m in enumerate(MONTHS_TH):
        ws.append([m, round(2_150_000 * random.uniform(0.9, 1.1)), round(410_000 * random.uniform(0.85, 1.15))])
    ws2 = wb.create_sheet("Water")
    ws2.append(["เดือน (2568)", "น้ำใช้ (ลบ.ม.)", "น้ำรีไซเคิล (ลบ.ม.)"])
    for m in MONTHS_TH:
        use = round(64_000 * random.uniform(0.85, 1.2))
        ws2.append([m, use, round(use * random.uniform(0.25, 0.4))])
    ws2.append(["หมายเหตุ", "เดือน ก.ค. มิเตอร์เสีย ใช้ค่าประมาณ", ""])
    base.mkdir(parents=True, exist_ok=True)
    wb.save(base / "site_A_energy_water_2568.xlsx")
    # Site C (coal plant): csv, English headers, electricity in MWh, coal in tonnes
    write_csv(base / "site_C_esg_2568.csv",
              ["month", "aux_power_MWh", "coal_t", "water_withdrawal_m3", "ash_t", "ash_reused_pct"],
              [[f"2025-{i+1:02d}", round(5_800 * random.uniform(0.9, 1.1)), round(61_000 * random.uniform(0.9, 1.1)),
                round(410_000 * random.uniform(0.9, 1.1)), round(5_900 * random.uniform(0.9, 1.1)),
                round(random.uniform(55, 80), 1)] for i in range(12)])
    # Site F (solar): xlsx, quarterly only (granularity mismatch)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ESG_Q"
    ws.append(["ไตรมาส", "ไฟฟ้าผลิตได้ (MWh)", "ไฟฟ้าใช้เอง (MWh)", "น้ำล้างแผง (ลบ.ม.)"])
    for q in range(1, 5):
        ws.append([f"Q{q}/2568", round(33_000 * random.uniform(0.85, 1.1)), round(260 * random.uniform(0.9, 1.1)),
                   round(190 * random.uniform(0.8, 1.3))])
    wb.save(base / "site_F_solar_2568.xlsx")
    # Site G (gas plant): csv with Thai headers, gas in MMBtu, one missing month
    rows = []
    for i, m in enumerate(MONTHS_TH):
        if i == 9:
            rows.append([m, "", "", "รอข้อมูลจาก Site"])
            continue
        rows.append([m, round(1_250_000 * random.uniform(0.9, 1.1)), round(150_000 * random.uniform(0.9, 1.1)), ""])
    write_csv(base / "site_G_utilities_2568.csv", ["เดือน", "ก๊าซธรรมชาติ (MMBtu)", "น้ำใช้ (ลบ.ม.)", "หมายเหตุ"], rows)
    write_csv(base / "emission_factors_SYNTHETIC.csv", ["source", "unit", "kgCO2e_per_unit", "note"], [
        ["grid_electricity", "kWh", 0.45, "ค่าสมมติเพื่อฝึก ไม่ใช่ค่าทางการ"],
        ["diesel", "litre", 2.7, "ค่าสมมติเพื่อฝึก ไม่ใช่ค่าทางการ"],
        ["coal", "tonne", 2_300, "ค่าสมมติเพื่อฝึก ไม่ใช่ค่าทางการ"],
        ["natural_gas", "MMBtu", 53, "ค่าสมมติเพื่อฝึก ไม่ใช่ค่าทางการ"],
    ])
    write_csv(base / "esg_targets_2568.csv", ["metric", "target", "unit", "owner"], [
        ["water_recycled_share_site_A", 30, "%", "ทีม ESG (สมมติ)"],
        ["ash_reused_site_C", 70, "%", "ทีม ESG (สมมติ)"],
        ["scope1_2_total", 1_950_000, "tCO2e", "ทีม ESG (สมมติ)"],
    ])


# ---------------------------------------------------------------- UC4
MEETINGS = {
    "2026-09-07_weekly_ops.md": """# ประชุม Ops ประจำสัปดาห์ (7 ก.ย. 2569) · ข้อมูลสมมติ
ผู้เข้าร่วม: คุณสมมติ วางแผน (ประธาน), คุณทดสอบ ซ่อมเก่ง, คุณตัวอย่าง จัดซื้อดี, คุณจำลอง ข้อมูลแม่น
- Site C แจ้งท่อไอน้ำรั่ว คุณทดสอบ จะประสานผู้รับเหมาให้เข้าซ่อมภายใน 18 ก.ย.
- ราคาอะไหล่สายพานขึ้น 4% คุณตัวอย่าง ไปเช็กเงื่อนไขปรับราคาในสัญญา CT-2569-014 ก่อนประชุมหน้า
- Dashboard ผลผลิตรายเดือนยังดึงข้อมูลช้า คุณจำลอง จะลองตั้ง scheduled task (ยังไม่กำหนดวัน)
- ทุกคนส่งตัวเลข Q3 ให้คุณสมมติ ภายใน 30 ก.ย.""",
    "2026-09-14_weekly_ops.txt": """ประชุม Ops 14 ก.ย. 2569 (บันทึกแบบย่อ, ข้อมูลสมมติ)
- ท่อไอน้ำ Site C: ผู้รับเหมาขอเลื่อนเป็น 25 ก.ย. -> คุณทดสอบ รับทราบ ขอให้แจ้ง HSE
- สัญญาสายพาน: ปรับได้ไม่เกิน 5%/ปี ต้องแจ้งล่วงหน้า 60 วัน -> ปิดเรื่อง
- คุณจำลอง ขอ 2 สัปดาห์ทำ scheduled task ต้นแบบ
- เพิ่ม: คุณตัวอย่าง ทำ shortlist ผู้รับจ้างล้างแผงรายใหม่ เผื่อสัญญา CT-2569-031 ไม่ต่อ ส่ง 5 ต.ค.""",
    "2026-09-21_esg_sync.vtt": """WEBVTT

00:00:05.000 --> 00:00:12.000
คุณสมมติ วางแผน: เริ่มเรื่อง ESG ปี 2568 ครับ ข้อมูลยังมาไม่ครบ

00:00:12.500 --> 00:00:21.000
คุณจำลอง ข้อมูลแม่น: Site G ยังไม่ส่งเดือนตุลาคม ผมจะตามให้ภายในวันศุกร์นี้

00:00:21.500 --> 00:00:30.000
คุณสมมติ วางแผน: ได้ครับ แล้วหน่วยของ Site A กับ Site C ไม่ตรงกัน ใครช่วยทำตารางแปลงหน่วย

00:00:30.500 --> 00:00:38.000
คุณทดสอบ ซ่อมเก่ง: ผมทำให้ครับ ขอถึงสิ้นเดือน

00:00:38.500 --> 00:00:45.000
คุณสมมติ วางแผน: โอเค ทั้งหมดนี้เป็นข้อมูลสมมติสำหรับการอบรมนะครับ""",
    "2026-09-28_weekly_ops.md": """# ประชุม Ops ประจำสัปดาห์ (28 ก.ย. 2569) · ข้อมูลสมมติ
- ท่อไอน้ำ Site C ซ่อมเสร็จ 25 ก.ย. ปิดเรื่อง
- คุณจำลอง: scheduled task ต้นแบบทำงานได้ แต่ต้องเปิดเครื่องทิ้งไว้ จะเสนอทางเลือกสัปดาห์หน้า
- ตัวเลข Q3 ยังขาด Site F -> คุณตัวอย่าง ตามให้ (ไม่ได้ระบุวัน)
- ความเสี่ยง: ถ้า shortlist ผู้รับจ้างล้างแผงช้า อาจไม่ทันแจ้งยกเลิก 30 วัน""",
}


def gen_uc4():
    base = ROOT / "uc4_meeting_actions"
    for fname, text in MEETINGS.items():
        write_text(base / "meetings" / fname, text)
    write_csv(base / "action_tracker_sharepoint_list_export.csv",
              ["ID", "Title", "Owner", "OwnerEmail", "DueDate", "Status", "Source"], [
        [101, "ส่งตัวเลข Q2 ให้ประธาน", "คุณจำลอง ข้อมูลแม่น", "jamlong.d@example.com", "2026-07-05", "Done", "ประชุม 28 มิ.ย."],
        [102, "อัปเดตแผนซ่อมบำรุง Site A", "คุณทดสอบ ซ่อมเก่ง", "thodsob.s@example.com", "2026-08-15", "In progress", "ประชุม 3 ส.ค."],
        [103, "ทบทวน SLA แคมป์ (CT-2569-022)", "คุณตัวอย่าง จัดซื้อดี", "tuayang.j@example.com", "2026-08-31", "Not started", "ประชุม 17 ส.ค."],
        [104, "ทดลอง Dashboard ใหม่", "คุณจำลอง ข้อมูลแม่น", "jamlong.d@example.com", "2026-09-10", "In progress", "ประชุม 24 ส.ค."],
        [105, "เตรียมข้อมูล ESG ครึ่งปี", "คุณสมมติ วางแผน", "sommut.w@example.com", "2026-09-30", "Not started", "ประชุม 31 ส.ค."],
    ])
    write_text(base / "owners.md", """# รายชื่อผู้รับผิดชอบ (สมมติทั้งหมด)
| ชื่อ | อีเมล | ทีม |
|---|---|---|
| คุณสมมติ วางแผน | sommut.w@example.com | Ops Planning |
| คุณทดสอบ ซ่อมเก่ง | thodsob.s@example.com | Maintenance |
| คุณตัวอย่าง จัดซื้อดี | tuayang.j@example.com | Procurement |
| คุณจำลอง ข้อมูลแม่น | jamlong.d@example.com | Data & Digital |""")


# ---------------------------------------------------------------- UC5
NEWS = [
    ("ราคาถ่านหินดัชนีสมมติ NEX-6000 ทรงตัว หลังจีนเพิ่มการผลิตในประเทศ", "coal"),
    ("ญี่ปุ่นนำเข้า LNG ลดลงเล็กน้อย จากอากาศอุ่นกว่าปกติ (สมมติ)", "lng"),
    ("กกพ.สมมติ เปิดรับฟังความเห็นอัตรารับซื้อไฟฟ้าพลังงานแสงอาทิตย์รอบใหม่", "power"),
    ("ราคาน้ำมันดิบอ้างอิงสมมติ ปรับขึ้น 3% จากความกังวลด้านอุปทาน", "oil"),
    ("อินโดนีเซียสมมติ พิจารณาภาษีส่งออกถ่านหินขั้นบันได", "coal"),
    ("ออสเตรเลียสมมติ ฝนตกหนักกระทบการขนส่งถ่านหินทางราง", "coal"),
    ("ค่าเงินบาทแข็งค่าแตะ 35.1 บาท/ดอลลาร์ (สมมติ)", "fx"),
    ("โครงการแบตเตอรี่กักเก็บพลังงานในภูมิภาคสมมติ เพิ่มขึ้น 40%", "power"),
    ("สหรัฐสมมติ ส่งออก LNG สูงสุดเป็นประวัติการณ์", "lng"),
    ("ตลาดคาร์บอนเครดิตสมมติ ราคาเฉลี่ยขึ้นเป็น 9 ดอลลาร์/ตัน", "carbon"),
    ("เวียดนามสมมติ เร่งเปิดประมูลโซลาร์หลังคา", "power"),
    ("ราคาถ่านหินดัชนีสมมติ NEX-6000 ลดลงสัปดาห์ที่ 2 จากสต็อกท่าเรือสูง", "coal"),
]


def gen_uc5():
    base = ROOT / "uc5_energy_market_brief"
    start = dt.date(2026, 8, 3)
    rows, news = [], []
    coal, lng, oil, fx = 112.0, 11.8, 78.0, 35.6
    for w in range(8):
        d = start + dt.timedelta(weeks=w)
        coal *= random.uniform(0.95, 1.05)
        lng *= random.uniform(0.94, 1.06)
        oil *= random.uniform(0.96, 1.04)
        fx += random.uniform(-0.25, 0.2)
        rows.append([d.isoformat(), round(coal, 1), round(lng, 2), round(oil, 1), round(fx, 2)])
    write_csv(base / "prices_weekly_SYNTHETIC.csv",
              ["week_start", "coal_index_NEX6000_usd_t", "lng_spot_usd_mmbtu", "crude_ref_usd_bbl", "thb_per_usd"], rows)
    for i, (h, tag) in enumerate(NEWS):
        d = start + dt.timedelta(days=4 * i + 2)
        news.append([d.isoformat(), h, tag, "ข่าวสมมติ (ไม่ใช่ข่าวจริง)", "https://news.example.com/item/" + str(1000 + i)])
    write_csv(base / "news_snippets_SYNTHETIC.csv", ["date", "headline_th", "topic", "note", "url"], news)
    write_text(base / "internal_view.md", f"""# มุมมองทีมกลยุทธ์ (สมมติ) · {NOTE}
- เราสนใจ 3 เรื่อง: ราคาถ่านหิน (ต้นทุน/รายได้), LNG (โรงไฟฟ้าก๊าซ), นโยบายพลังงานหมุนเวียนในภูมิภาค
- ผู้อ่าน: ผู้บริหารระดับกลาง อ่านบนมือถือ ไม่เกิน 1 หน้า
- ห้ามคาดการณ์ราคาเป็นตัวเลข ให้บอกแนวโน้ม + เหตุผล + แหล่งที่มา
- ส่งทุกเช้าวันจันทร์ (ทีม M6 ตั้งเป็น scheduled task ได้)""")


def make_zips():
    out = ROOT / "out"
    out.mkdir(exist_ok=True)
    for d in sorted(ROOT.glob("uc*_*")):
        with zipfile.ZipFile(out / f"M7_{d.name}.zip", "w", zipfile.ZIP_DEFLATED) as z:
            for f in sorted(d.rglob("*")):
                if f.is_file():
                    z.write(f, f.relative_to(ROOT))


if __name__ == "__main__":
    gen_uc1()
    gen_uc2()
    gen_uc3()
    gen_uc4()
    gen_uc5()
    make_zips()
    print("M7 data pack written to", ROOT)
